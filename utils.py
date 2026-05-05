import io
import pandas as pd
import streamlit as st

# --- CONSTANTS ---
STANDARD_CAPACITY = 152
MAX_HOURS_PER_PROGRAM = int(STANDARD_CAPACITY * 1.2)

UTILIZATION_WARN_THRESHOLD = 80
UTILIZATION_OVER_THRESHOLD = 100
UNDO_STACK_LIMIT = 10
NAME_TRUNCATION_LENGTH = 12
JUMP_BUTTON_COLUMNS = 5
JUMP_BUTTON_DISPLAY_LIMIT = 10

RATE_CARD = {
    "ACP": 37,
    "CP": 54,
    "SCP": 54,
    "LCP": 89,
    "ACE": 89,
    "CE": 89,
    "SCE": 119,
    "R+I I": 44,
    "R+I II": 56,
    "R+I III": 89,
    "R+I IV": 135
}

TEAM_ROLES = ['ACP', 'CP', 'SCP', 'LCP', 'ACE', 'CE', 'SCE']


def get_program_columns(df: pd.DataFrame) -> list[str]:
    exclude = {'Capacity', 'Current Hours to Target'}
    return [c for c in df.select_dtypes(include=['number']).columns if c not in exclude]


def find_column(df: pd.DataFrame, candidates: list[str]) -> str | None:
    """Finds a column name from a list of candidates (case-insensitive)."""
    df_cols_clean = [str(col).strip().lower() for col in df.columns]
    for c in candidates:
        c_clean = c.strip().lower()
        if c_clean in df_cols_clean:
            return df.columns[df_cols_clean.index(c_clean)]
    return None


def process_uploaded_file(file) -> tuple[pd.DataFrame, dict[str, float]]:
    """Parses, Pivots, and Normalizes incoming CSV data."""
    try:
        file.seek(0)
        df = pd.read_csv(file)
        df.columns = df.columns.str.strip()

        mrr_col = find_column(df, ['Program MRR', 'MRR', 'Revenue'])
        prog_col_raw = find_column(df, ['Program Name', 'Program', 'Client'])

        new_mrr_map = {}
        if mrr_col and prog_col_raw:
            try:
                if df[mrr_col].dtype == 'object':
                    df[mrr_col] = df[mrr_col].astype(str).str.replace('$', '').str.replace(',', '')
                df[mrr_col] = pd.to_numeric(df[mrr_col], errors='coerce').fillna(0)
                temp_map = df.groupby(prog_col_raw)[mrr_col].max().to_dict()
                new_mrr_map.update(temp_map)
            except Exception as e:
                st.warning(f"Could not parse MRR values: {e}")
            df = df.drop(columns=[mrr_col])

        ct_col = find_column(df, ['CT Name', 'Employee Name', 'Employee'])
        prog_col = find_column(df, ['Program Name', 'Program', 'Client'])
        role_col = find_column(df, ['Account Role', 'Role'])

        if ct_col and prog_col:
            hour_col = find_column(df, ['Allocated Monthly Hours', 'Allocated Hours', 'Hours'])
            if not hour_col:
                found = [c for c in df.columns if 'Allocated' in c]
                if found: hour_col = found[0]

            if hour_col:
                df[hour_col] = pd.to_numeric(df[hour_col], errors='coerce').fillna(0)
                pivot_df = df.pivot_table(index=ct_col, columns=prog_col, values=hour_col, aggfunc='sum').fillna(0)

                if role_col:
                    roles = df[[ct_col, role_col]].drop_duplicates(subset=ct_col).groupby(ct_col).first()
                    final_df = roles.join(pivot_df).reset_index()
                    final_df = final_df.rename(columns={ct_col: 'Employee', role_col: 'Role'})
                else:
                    final_df = pivot_df.reset_index().rename(columns={ct_col: 'Employee'})

                return final_df, new_mrr_map

        if ct_col and ct_col != 'Employee':
            df = df.rename(columns={ct_col: 'Employee'})

        return df, new_mrr_map

    except Exception as e:
        st.error(f"Error processing file: {e}")
        return pd.DataFrame(), {}


def recalculate_utilization(df: pd.DataFrame) -> pd.DataFrame:
    """Updates the 'Current Hours to Target' column based on allocated hours."""
    if df.empty: return df

    prog_cols = get_program_columns(df)
    total_hours = df[prog_cols].sum(axis=1)

    if 'Capacity' not in df.columns:
        col_idx = 1 if 'Role' in df.columns else 0
        df.insert(col_idx, 'Capacity', STANDARD_CAPACITY)

    util = df.apply(lambda x: (total_hours[x.name] / x['Capacity'] * 100) if x['Capacity'] > 0 else 0, axis=1)
    df['Current Hours to Target'] = util.round(0).astype(int)

    return df


def get_rate(role_name: str | None) -> int:
    if not role_name: return 0
    role_clean = str(role_name).strip().upper()
    return RATE_CARD.get(role_clean, 0)


def calculate_margin(df: pd.DataFrame, program_mrr_dict: dict[str, float]) -> dict[str, dict]:
    """Calculates Extended Cost and Margin % for all programs."""
    if df.empty: return {}
    prog_cols = get_program_columns(df)
    margin_data = {}

    rates = df['Role'].map(get_rate).fillna(0) if 'Role' in df.columns else pd.Series(0, index=df.index)
    program_costs = {prog: (df[prog] * rates).sum() for prog in prog_cols}

    for prog, cost in program_costs.items():
        mrr = program_mrr_dict.get(prog, 0)
        if mrr > 0:
            margin_pct = ((mrr - cost) / mrr) * 100
        else:
            margin_pct = -100.0 if cost > 0 else 0.0
        margin_data[prog] = {'mrr': mrr, 'cost': cost, 'margin_pct': margin_pct}
    return margin_data


def calculate_group_utilization(df: pd.DataFrame, role_list: list[str]) -> tuple[float, float, float]:
    """Total Hours / (Count of Employees * STANDARD_CAPACITY), filtered by role."""
    if 'Role' not in df.columns or df.empty:
        return 0, 0, 0

    mask = df['Role'].astype(str).str.upper().isin([r.upper() for r in role_list])
    role_df = df[mask]

    if role_df.empty:
        return 0, 0, 0

    prog_cols = get_program_columns(role_df)
    total_allocated_hours = role_df[prog_cols].sum().sum()
    total_capacity = len(role_df) * STANDARD_CAPACITY

    utilization_pct = 0
    if total_capacity > 0:
        utilization_pct = (total_allocated_hours / total_capacity) * 100

    return utilization_pct, total_allocated_hours, total_capacity


def export_to_excel(current_df: pd.DataFrame, original_df: pd.DataFrame | None, program_mrr: dict[str, float]) -> bytes:
    """Exports the current DataFrame to Excel with changed cells highlighted."""
    from openpyxl.styles import PatternFill

    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    green_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")

    output = io.BytesIO()
    export_df = current_df.copy()
    export_df.index.name = 'Employee'

    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        export_df.to_excel(writer, sheet_name='Staffing')
        ws = writer.sheets['Staffing']

        if original_df is not None and not original_df.empty:
            for row_idx, emp_name in enumerate(export_df.index, start=2):
                is_new_row = emp_name not in original_df.index
                for col_idx, col_name in enumerate(export_df.columns, start=2):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    is_new_col = col_name not in original_df.columns

                    if is_new_row or is_new_col:
                        cell.fill = green_fill
                    elif emp_name in original_df.index and col_name in original_df.columns:
                        current_val = export_df.at[emp_name, col_name]
                        original_val = original_df.at[emp_name, col_name]
                        try:
                            if current_val != original_val:
                                cell.fill = yellow_fill
                        except (TypeError, ValueError):
                            pass

    return output.getvalue()
